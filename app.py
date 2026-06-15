"""
Analisador de Currículos com IA (OpenAI)
Flask app pronto para deploy no Render.com
"""
import os
import io
import json
import uuid
import sqlite3
from datetime import datetime
from pathlib import Path

from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for
from werkzeug.utils import secure_filename
from openai import OpenAI
from pypdf import PdfReader
from docx import Document
from openpyxl import Workbook

# ---------- Config ----------
BASE_DIR = Path(__file__).parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
DB_PATH = BASE_DIR / "data.db"
ALLOWED_EXT = {"pdf", "doc", "docx", "txt"}
MAX_FILE_MB = 10

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50MB total

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
client = OpenAI(api_key=OPENAI_API_KEY) if OPENAI_API_KEY else None


# ---------- DB ----------
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS jobs (
            id TEXT PRIMARY KEY,
            title TEXT, company TEXT, description TEXT,
            required TEXT, desired TEXT, education TEXT, experience TEXT,
            hard_skills TEXT, soft_skills TEXT, notes TEXT,
            weights TEXT, created_at TEXT
        );
        CREATE TABLE IF NOT EXISTS candidates (
            id TEXT PRIMARY KEY,
            job_id TEXT,
            filename TEXT,
            name TEXT, location TEXT, whatsapp TEXT, linkedin TEXT,
            score INTEGER, summary TEXT, details TEXT,
            stage TEXT DEFAULT 'Novo',
            created_at TEXT
        );
        """)


init_db()


# ---------- Helpers ----------
def allowed(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXT


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext in (".docx", ".doc"):
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        if ext == ".txt":
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return f"[Erro lendo arquivo: {e}]"
    return ""


ANALYSIS_PROMPT = """Você é um avaliador técnico de RH altamente experiente.
Avalie o currículo abaixo em relação à vaga descrita.
IMPORTANTE: não considere gênero, idade, nacionalidade ou foto para a pontuação.

VAGA:
Título: {title}
Empresa/Setor: {company}
Descrição: {description}
Requisitos obrigatórios: {required}
Requisitos desejáveis: {desired}
Escolaridade mínima: {education}
Experiência mínima: {experience}
Hard skills: {hard_skills}
Soft skills: {soft_skills}
Observações: {notes}

PESOS (soma = 100):
{weights}

CURRÍCULO:
{cv_text}

Responda APENAS em JSON válido com este schema exato:
{{
  "name": "Nome completo do candidato",
  "location": "Cidade - UF se identificável, ou ''",
  "whatsapp": "telefone com DDD se identificável, ou ''",
  "linkedin": "url do linkedin se houver, ou ''",
  "score": <inteiro 0-100>,
  "summary": "1-2 frases sobre aderência",
  "scores_breakdown": {{
    "experiencia": <0-100>,
    "comp_tecnicas": <0-100>,
    "formacao": <0-100>,
    "certificacoes": <0-100>,
    "resultados": <0-100>,
    "comp_comportamentais": <0-100>
  }},
  "strengths": ["..."],
  "gaps": ["..."]
}}
"""


def analyze_with_ai(job: dict, cv_text: str) -> dict:
    if not client:
        # Fallback determinístico para teste sem chave
        return {
            "name": "Candidato Demo",
            "location": "",
            "whatsapp": "",
            "linkedin": "",
            "score": 50,
            "summary": "OPENAI_API_KEY não configurada — análise simulada.",
            "scores_breakdown": {},
            "strengths": [],
            "gaps": ["Configure OPENAI_API_KEY"],
        }
    prompt = ANALYSIS_PROMPT.format(
        cv_text=cv_text[:15000],
        weights=json.dumps(job.get("weights", {}), ensure_ascii=False),
        **{k: job.get(k, "") for k in [
            "title", "company", "description", "required", "desired",
            "education", "experience", "hard_skills", "soft_skills", "notes"
        ]},
    )
    resp = client.chat.completions.create(
        model=OPENAI_MODEL,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
        temperature=0.2,
    )
    return json.loads(resp.choices[0].message.content)


# ---------- Routes ----------
@app.route("/")
def index():
    return redirect(url_for("new_analysis"))


@app.route("/nova")
def new_analysis():
    return render_template("new.html")


@app.route("/api/analyze", methods=["POST"])
def api_analyze():
    form = request.form
    files = request.files.getlist("cvs")
    if not files:
        return jsonify({"error": "Nenhum currículo enviado"}), 400

    try:
        weights = {
            "experiencia": int(form.get("w_exp", 25)),
            "comp_tecnicas": int(form.get("w_tech", 25)),
            "formacao": int(form.get("w_form", 15)),
            "certificacoes": int(form.get("w_cert", 10)),
            "resultados": int(form.get("w_result", 15)),
            "comp_comportamentais": int(form.get("w_soft", 10)),
        }
    except ValueError:
        return jsonify({"error": "Pesos inválidos"}), 400

    if sum(weights.values()) != 100:
        return jsonify({"error": "A soma dos pesos deve ser 100"}), 400

    job = {
        "title": form.get("title", "").strip(),
        "company": form.get("company", "").strip(),
        "description": form.get("description", ""),
        "required": form.get("required", ""),
        "desired": form.get("desired", ""),
        "education": form.get("education", ""),
        "experience": form.get("experience", ""),
        "hard_skills": form.get("hard_skills", ""),
        "soft_skills": form.get("soft_skills", ""),
        "notes": form.get("notes", ""),
        "weights": weights,
    }
    if not job["title"] or not job["company"]:
        return jsonify({"error": "Título da vaga e Empresa são obrigatórios"}), 400

    job_id = "vaga_" + uuid.uuid4().hex[:16]
    now = datetime.utcnow().isoformat()

    with db() as c:
        c.execute("""INSERT INTO jobs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
            job_id, job["title"], job["company"], job["description"],
            job["required"], job["desired"], job["education"], job["experience"],
            job["hard_skills"], job["soft_skills"], job["notes"],
            json.dumps(weights), now,
        ))

    results = []
    for f in files:
        if not f or not f.filename or not allowed(f.filename):
            continue
        fname = secure_filename(f.filename)
        dest = UPLOAD_DIR / f"{uuid.uuid4().hex}_{fname}"
        f.save(dest)
        text = extract_text(dest)
        if not text.strip():
            continue
        try:
            r = analyze_with_ai(job, text)
        except Exception as e:
            r = {"name": fname, "score": 0, "summary": f"Erro IA: {e}",
                 "location": "", "whatsapp": "", "linkedin": "",
                 "scores_breakdown": {}, "strengths": [], "gaps": []}

        cid = "cand_" + uuid.uuid4().hex[:16]
        with db() as c:
            c.execute("""INSERT INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""", (
                cid, job_id, fname,
                r.get("name") or fname, r.get("location", ""),
                r.get("whatsapp", ""), r.get("linkedin", ""),
                int(r.get("score", 0)), r.get("summary", ""),
                json.dumps(r, ensure_ascii=False), "Novo", now,
            ))
        results.append({"id": cid, "name": r.get("name") or fname,
                        "score": r.get("score", 0)})

    return jsonify({"job_id": job_id, "count": len(results), "results": results})


@app.route("/resultados")
def results_page():
    return render_template("results.html")


@app.route("/api/candidates")
def api_candidates():
    job_filter = request.args.get("job", "")
    level = request.args.get("level", "")
    stage = request.args.get("stage", "")
    q = request.args.get("q", "").lower()

    with db() as c:
        rows = c.execute("""
            SELECT c.*, j.title AS job_title, j.id AS job_code
            FROM candidates c LEFT JOIN jobs j ON c.job_id = j.id
            ORDER BY c.created_at DESC
        """).fetchall()
        jobs = c.execute("SELECT id, title FROM jobs ORDER BY created_at DESC").fetchall()

    def lvl(score):
        if score >= 80: return "Alta aderência"
        if score >= 60: return "Boa aderência"
        if score >= 40: return "Aderência parcial"
        return "Baixa aderência"

    out = []
    for r in rows:
        item = dict(r)
        item["level"] = lvl(item["score"])
        if job_filter and item["job_id"] != job_filter: continue
        if stage and item["stage"] != stage: continue
        if level and item["level"] != level: continue
        if q and q not in (item["name"] or "").lower() and q not in (item["job_title"] or "").lower():
            continue
        out.append(item)

    return jsonify({
        "candidates": out,
        "jobs": [dict(j) for j in jobs],
    })


@app.route("/api/candidates/<cid>/stage", methods=["POST"])
def update_stage(cid):
    stage = request.json.get("stage", "Novo")
    with db() as c:
        c.execute("UPDATE candidates SET stage=? WHERE id=?", (stage, cid))
    return jsonify({"ok": True})


@app.route("/api/candidates/<cid>", methods=["DELETE"])
def delete_candidate(cid):
    with db() as c:
        c.execute("DELETE FROM candidates WHERE id=?", (cid,))
    return jsonify({"ok": True})


@app.route("/api/candidates/<cid>/detail")
def candidate_detail(cid):
    with db() as c:
        r = c.execute("""SELECT c.*, j.title AS job_title FROM candidates c
                         LEFT JOIN jobs j ON c.job_id=j.id WHERE c.id=?""", (cid,)).fetchone()
    if not r: return jsonify({"error": "not found"}), 404
    item = dict(r)
    try: item["details"] = json.loads(item["details"])
    except: pass
    return jsonify(item)


@app.route("/export.xlsx")
def export_xlsx():
    with db() as c:
        rows = c.execute("""SELECT c.*, j.title AS job_title, j.id AS job_code
                            FROM candidates c LEFT JOIN jobs j ON c.job_id=j.id
                            ORDER BY c.score DESC""").fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidatos"
    ws.append(["Nome", "Localização", "WhatsApp", "LinkedIn", "Vaga", "Código Vaga",
               "Compatibilidade (%)", "Etapa", "Resumo IA", "Data"])
    for r in rows:
        ws.append([r["name"], r["location"], r["whatsapp"], r["linkedin"],
                   r["job_title"] or "", r["job_code"] or "",
                   r["score"], r["stage"], r["summary"], r["created_at"]])
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=f"candidatos_{datetime.now():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/healthz")
def health():
    return {"status": "ok", "openai_configured": bool(OPENAI_API_KEY)}


if __name__ == "__main__":
    port = int(os.getenv("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
